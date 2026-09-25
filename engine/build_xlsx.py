"""Génère le dossier de paramétrage GTA #Dièse (classeur Excel).

python3 build_xlsx.py --input export_regles.xlsx [--output dossier.xlsx]
                     [--daytypes export_absences.xlsx] [--client "Nom"]
                     [--ref famille=fichier.xlsx] [--langue fr|en]
                     [--contrats contrats|non] [--sample 12,45,78] [--diagnostic]
"""
import argparse
import datetime as dt
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from dossier_common import (ExportError, _as_int, load_rules, load_daytypes, effective_order, clean_label,
                            clean_code,
                            contract_index, diagnose, format_diagnosis, load_referentials,
                            load_daytype_details, daytype_usage)
from humanizer import Resolver, FormulaTranslator, describe, HANDLERS

VERSION = '2.2'
FONT = 'Arial'
RED = 'B3202C'          # rouge IT4culture : bandeaux et en-têtes
RED_DARK = '7E1620'     # rouge sombre : titres de section
RED_PALE = 'F8E7E9'     # rouge très pâle : compteurs, alternance de la notice
INK = '2B2B2B'          # gris anthracite : texte courant
GREY = 'F3F4F6'         # gris clair : lignes alternées du catalogue
RULE_LINE = 'E4E6EA'    # gris : filets

COLUMNS = {
    'fr': [('Ordre', 8), ('#', 7), ('Libellé', 40), ('Libellé court', 22), ('Code', 12),
           ('Ce que fait la règle', 100), ('Compteur', 10), ('Période', 18)],
    'en': [('Order', 8), ('#', 7), ('Name', 40), ('Short label', 22), ('Code', 12),
           ('What the rule calculates', 100), ('Counter', 10), ('Period', 18)],
}

GLOSSARY_COLUMNS = {
    'fr': [('Intitulé public', 34), ('Intitulé court', 20), ('#', 7), ('Couleur', 11),
           ('Nature', 30), ('Durée par défaut', 16), ('Actif', 9), ('Règles concernées', 17)],
    'en': [('Public label', 34), ('Short label', 20), ('#', 7), ('Colour', 11),
           ('Nature', 30), ('Default duration', 16), ('Active', 9), ('Rules using it', 17)],
}

UI = {
    'title': ('DOSSIER DE PARAMÉTRAGE GTA', 'GTA CONFIGURATION DOSSIER'),
    'subtitle': ('Environnement #Dièse', '#Dièse environment'),
    'module': ('Module Gestion des Temps et des Activités  ·  généré le {date}',
               'Time and Activity Management module  ·  generated on {date}'),
    'section': ('À QUOI SERT CE DOCUMENT', 'WHAT THIS DOCUMENT IS FOR'),
    'intro': (
        "Ce dossier décrit, règle par règle, la façon dont votre environnement #Dièse calcule "
        "les temps et les activités. Chaque ligne du catalogue correspond à une règle réellement "
        "paramétrée, accompagnée d'une explication en français de ce qu'elle calcule.\n\n"
        "Les règles y figurent dans l'ordre où #Dièse les exécute : une règle peut réutiliser le "
        "résultat de celles qui la précèdent. Un renvoi noté « Libellé (#12) » désigne une autre "
        "règle du catalogue ; un élément noté ‹ entre chevrons › n'a pas pu être nommé faute du "
        "référentiel correspondant, et sa référence d'origine est conservée telle quelle.\n\n"
        "Ce document s'adresse aux équipes RH, paie et planification : aucune connaissance "
        "technique n'est nécessaire pour le lire.",
        "This dossier describes, rule by rule, how your #Dièse environment calculates time and "
        "activities. Each row of the catalogue is a rule that is actually configured, with an "
        "explanation in plain English of what it calculates.\n\n"
        "Rules appear in the order #Dièse runs them: a rule may reuse the result of those before "
        "it. A reference written “Name (#12)” points to another rule in this catalogue; an item "
        "written ‹ in angle brackets › could not be named because the matching reference data was "
        "not supplied, so its original reference is kept as is.\n\n"
        "Labels of rules, day types and contract types come from your own environment and are "
        "therefore left in their original language.\n\n"
        "This document is written for HR, payroll and scheduling teams: no technical knowledge "
        "is required."),
    'footer': ('IT4culture  ·  dossier de paramétrage GTA #Dièse  ·  générateur v{version}',
               'IT4culture  ·  #Dièse GTA configuration dossier  ·  generator v{version}'),
    'catalog': ('Catalogue des règles', 'Rule catalogue'),
    'columns_section': ('LES COLONNES DU CATALOGUE', 'CATALOGUE COLUMNS'),
    'glossary_section': ('LES COLONNES DU LEXIQUE DES TYPES DE JOUR', 'DAY TYPE GLOSSARY COLUMNS'),
    'contracts_key': ('Colonnes de contrat', 'Contract columns'),
    'filter_tip': (
        "Astuce : chaque en-tête porte une flèche de filtre. Pour ne voir que les règles d'un "
        "type de contrat, filtrez sa colonne sur « X » ; pour ne voir que les compteurs, filtrez "
        "la colonne Compteur sur « Oui ». Les premières colonnes restent visibles quand on fait "
        "défiler le tableau vers la droite.",
        "Tip: every header has a filter arrow. To see only the rules of one contract type, "
        "filter its column on “X”; to see only counters, filter the Counter column on “Yes”. "
        "The first columns stay visible when you scroll the table to the right."),

    'glossary': ('Lexique des types de jour', 'Day type glossary'),
    'readme': ('Lisez-moi', 'Read me'),
    'yes': ('Oui', 'Yes'),
    'no': ('Non', 'No'),
}


# Mode d'emploi des colonnes, dans l'ordre du catalogue. Les clés reprennent
# exactement les en-têtes de COLUMNS : un en-tête renommé doit l'être ici aussi.
COLUMN_GUIDE = {
    'fr': [
        ('Ordre', "Position de la règle dans la chaîne de calcul. #Dièse exécute les règles de "
                  "la plus petite valeur à la plus grande ; à ordre égal, dans l'ordre des "
                  "numéros. Une règle peut donc réutiliser le résultat de celles qui la "
                  "précèdent."),
        ('#', "Numéro unique de la règle dans #Dièse. C'est lui qui figure entre parenthèses "
              "dans les renvois d'une règle à une autre : « Libellé (#12) »."),
        ('Libellé', "Nom complet de la règle, tel qu'il est saisi dans le paramétrage #Dièse."),
        ('Libellé court', "Nom abrégé affiché dans les écrans et les récapitulatifs de #Dièse. "
                          "C'est le repère le plus sûr pour retrouver la règle à l'écran."),
        ('Code', "Code attribué à la règle, par exemple pour un export vers la paie. Vide "
                 "lorsqu'aucun code n'est paramétré."),
        ('Ce que fait la règle', "Explication en langage courant de ce que la règle calcule. "
                                 "Les formules les plus longues sont écrites en abrégé "
                                 "(« → » pour « alors », symboles de comparaison, libellés "
                                 "courts), sans rien omettre."),
        ('Compteur', "« Oui » lorsque le résultat de la règle se cumule sur une période, comme "
                     "un solde, au lieu d'être calculé jour par jour."),
        ('Période', "Pour un compteur, la période sur laquelle le résultat se cumule. « — » "
                    "pour les autres règles."),
    ],
    'en': [
        ('Order', "Position of the rule in the calculation chain. #Dièse runs rules from the "
                  "lowest value to the highest; rules with the same value run in order of "
                  "their number. A rule can therefore reuse the result of the rules before it."),
        ('#', "Unique number of the rule in #Dièse. It is the number shown in brackets when a "
              "rule refers to another: “Name (#12)”."),
        ('Name', "Full name of the rule, as entered in the #Dièse configuration."),
        ('Short label', "Short name displayed in #Dièse screens and summaries. It is the most "
                        "reliable way to find the rule on screen."),
        ('Code', "Code assigned to the rule, for instance for a payroll export. Empty when no "
                 "code is configured."),
        ('What the rule calculates', "Plain-English explanation of what the rule calculates. "
                                     "The longest formulas are written in shorthand (“→” for "
                                     "“then”, comparison symbols, short labels) without "
                                     "leaving anything out."),
        ('Counter', "“Yes” when the result of the rule builds up over a period, like a "
                    "balance, instead of being calculated day by day."),
        ('Period', "For a counter, the period over which the result builds up. “—” for other "
                   "rules."),
    ],
}

CONTRACT_GUIDE = (
    "Une colonne par type de contrat de l'environnement ({examples}). Une croix « X » indique "
    "que la règle s'applique aux salariés relevant de ce type de contrat. Une ligne sans aucune "
    "croix est une règle affectée à aucun contrat : elle reste parfois utilisée par d'autres "
    "règles, c'est pourquoi elle figure au catalogue.",
    "One column per contract type in the environment ({examples}). An “X” means the rule "
    "applies to employees under that contract type. A row with no “X” at all is a rule "
    "assigned to no contract: other rules sometimes still use it, which is why it is listed.")

GLOSSARY_GUIDE = {
    'fr': [
        ('Intitulé public / court', "Nom du type de jour tel qu'il apparaît dans les plannings. "
                                    "Le lexique est trié par ordre alphabétique."),
        ('#', "Numéro du type de jour, repris dans le catalogue lorsqu'il n'a pas pu être "
              "nommé."),
        ('Couleur', "Couleur du type de jour dans les plannings #Dièse."),
        ('Nature', "Traitement du type de jour dans le calcul des temps : temps de travail "
                   "effectif, non effectif ou absence."),
        ('Durée par défaut', "Durée retenue pour une journée de ce type lorsqu'aucun horaire "
                             "n'est saisi."),
        ('Actif', "« Non » pour un type de jour qui ne peut plus être posé dans les plannings."),
        ('Règles concernées', "Nombre de règles du catalogue qui s'appuient sur ce type de "
                              "jour. Zéro signale un type de jour qu'aucune règle n'utilise."),
    ],
    'en': [
        ('Public / short label', "Name of the day type as shown in schedules. The glossary is "
                                 "sorted alphabetically."),
        ('#', "Number of the day type, shown in the catalogue when it could not be named."),
        ('Colour', "Colour of the day type in #Dièse schedules."),
        ('Nature', "How the day type is treated in time calculation: effective working time, "
                   "non-effective working time or absence."),
        ('Default duration', "Duration used for a day of this type when no working hours are "
                             "entered."),
        ('Active', "“No” for a day type that can no longer be used in schedules."),
        ('Rules using it', "Number of rules in the catalogue that rely on this day type. Zero "
                           "flags a day type that no rule uses."),
    ],
}


def ui(lang, key, **kw):
    text = UI[key][0 if lang == 'fr' else 1]
    return text.format(**kw) if kw else text

thin = Side(style='thin', color=RULE_LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def default_output(rules_path, client=None, lang='fr'):
    """Dossier-parametrage-GTA_<client>_<date>.xlsx, à côté de l'export."""
    name = client or 'environnement'
    slug = re.sub(r'[^\w-]+', '-', name, flags=re.UNICODE).strip('-') or 'environnement'
    stem = 'Dossier-parametrage-GTA' if lang == 'fr' else 'GTA-configuration-dossier'
    return os.path.join(os.path.dirname(os.path.abspath(str(rules_path))),
                        '%s_%s_%s.xlsx' % (stem, slug, dt.date.today().isoformat()))


def translators(df, daytypes, referentials, lang):
    """(resolver, traducteur, traducteur raccourci) pour une langue.

    Le dictionnaire de résolution est construit sur TOUT l'export, jamais sur
    un échantillon : sinon les renvois entre règles cessent de se résoudre."""
    labels = {int(r['#']): r['Libellé'] for r in df}
    shorts = {int(r['#']): clean_label(r['Libellé court'], inline=True) for r in df}
    referentials = referentials or {}
    resolver = Resolver(labels, daytypes, referentials=referentials, lang=lang)
    translator = FormulaTranslator(resolver)
    compact = FormulaTranslator(Resolver(labels, daytypes, referentials=referentials,
                                         short_labels=shorts, lang=lang), compact=True)
    return resolver, translator, compact


def build(rules_path, out_path=None, daytypes_path=None, client=None, sample=None,
          contracts='contrats', lang='fr', referentials=None):
    """Écrit le classeur et renvoie son chemin. L'onglet lexique n'apparaît que
    si l'export des types de jour a été fourni : son absence n'est jamais bloquante."""
    out_path = out_path or default_output(rules_path, client, lang)
    df = load_rules(rules_path)
    daytypes = load_daytypes(daytypes_path) if daytypes_path else {}
    referentials = referentials or {}
    resolver, translator, compact = translators(df, daytypes, referentials, lang)

    columns, per_rule = ([], {})
    if contracts == 'contrats':
        columns, per_rule = contract_index(df)

    if sample:
        absent = sorted(set(sample) - {r['#'] for r in df})
        if absent:
            print("Échantillon : règle(s) absente(s) de l'export, ignorée(s) : %s"
                  % ', '.join('#%d' % i for i in absent))
    work = df.derive([r for r in df if r['#'] in set(sample)]) if sample else df
    work = effective_order(work)

    wb = Workbook()
    details = load_daytype_details(daytypes_path, lang) if daytypes_path else []
    _sheet_readme(wb.active, client, lang, columns, bool(details))
    _sheet_catalog(wb.create_sheet(ui(lang, 'catalog')), work, resolver, translator,
                   compact, columns, per_rule, lang)
    if details:
        _sheet_glossary(wb.create_sheet(ui(lang, 'glossary')), details, daytype_usage(df), lang)
    for ws in wb.worksheets:
        _print_setup(ws, landscape=(ws is not wb.worksheets[0]))
    wb.save(out_path)
    return out_path


def _print_setup(ws, landscape):
    """Impression sur la largeur d'une page, en-têtes répétés : le dossier est
    souvent imprimé ou converti en PDF avant d'être transmis."""
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if landscape:
        ws.print_title_rows = '1:1'


# --- Lisez-moi --------------------------------------------------------------

def _height(text, chars_per_line=96, line=13):
    """Hauteur de ligne pour un texte renvoyé à la ligne dans une cellule
    fusionnée : Excel ne l'ajuste pas seul."""
    lines = sum(max(1, -(-len(par) // chars_per_line)) for par in str(text).split('\n'))
    return lines * line + 5


def _sheet_readme(ws, client, lang='fr', columns=None, has_glossary=False):
    ws.title = ui(lang, 'readme')
    ws.sheet_view.showGridLines = False
    widths = {'A': 2.5, 'B': 30, 'C': 14, 'D': 14, 'E': 62}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    LAST = 'E'

    def band(row, height, color):
        for col in range(1, 6):
            ws.cell(row=row, column=col).fill = PatternFill('solid', fgColor=color)
        ws.row_dimensions[row].height = height

    def section(row, text):
        c = ws.cell(row=row, column=2, value=text.upper())
        c.font = Font(name=FONT, size=11, bold=True, color=RED_DARK)
        ws.merge_cells('B%d:%s%d' % (row, LAST, row))
        ws.row_dimensions[row].height = 22
        edge = Side(style='medium', color=RED)
        for col in range(2, 6):
            cell = ws.cell(row=row, column=col)
            cell.border = Border(bottom=edge)
        return row + 1

    def item(row, key, value, shade=False):
        a = ws.cell(row=row, column=2, value=key)
        a.font = Font(name=FONT, size=10, bold=True, color=INK)
        a.alignment = Alignment(vertical='top', wrap_text=True)
        b = ws.cell(row=row, column=3, value=value)
        b.font = Font(name=FONT, size=10, color=INK)
        b.alignment = Alignment(vertical='top', wrap_text=True)
        ws.merge_cells('C%d:%s%d' % (row, LAST, row))
        ws.row_dimensions[row].height = _height(value)
        if shade:
            for col in range(2, 6):
                ws.cell(row=row, column=col).fill = PatternFill('solid', fgColor=RED_PALE)
        return row + 1

    # ---- bandeau de titre
    for row in (1, 2, 3, 4, 5):
        band(row, 10 if row in (1, 5) else 26, RED)
    t = ws.cell(row=2, column=2, value=ui(lang, 'title'))
    t.font = Font(name=FONT, size=20, bold=True, color='FFFFFF')
    t.alignment = Alignment(vertical='center')
    ws.merge_cells('B2:%s2' % LAST)
    s = ws.cell(row=3, column=2, value=(client or ui(lang, 'subtitle')))
    s.font = Font(name=FONT, size=12, color='FFFFFF')
    s.alignment = Alignment(vertical='center')
    ws.merge_cells('B3:%s3' % LAST)
    d = ws.cell(row=4, column=2,
                value=ui(lang, 'module',
                         date=dt.date.today().strftime('%d/%m/%Y' if lang == 'fr' else '%Y-%m-%d')))
    d.font = Font(name=FONT, size=9, color='F3C9CD')
    d.alignment = Alignment(vertical='center')
    ws.merge_cells('B4:%s4' % LAST)

    r = 7
    r = section(r, ui(lang, 'section'))
    r += 1
    intro = ws.cell(row=r, column=2, value=ui(lang, 'intro'))
    intro.font = Font(name=FONT, size=10, color=INK)
    intro.alignment = Alignment(vertical='top', wrap_text=True)
    ws.merge_cells('B%d:%s%d' % (r, LAST, r))
    ws.row_dimensions[r].height = _height(ui(lang, "intro"), chars_per_line=125)
    r += 2

    # ---- mode d'emploi des colonnes
    r += 1
    r = section(r, ui(lang, 'columns_section'))
    ws.row_dimensions[r].height = 6
    r += 1
    guide = list(COLUMN_GUIDE[lang])
    if columns:
        names = [title for title, _k in columns]
        examples = ', '.join(names[:3]) + ('…' if len(names) > 3 else '')
        guide.append((ui(lang, 'contracts_key'),
                      CONTRACT_GUIDE[0 if lang == 'fr' else 1].format(examples=examples)))
    for i, (key, text) in enumerate(guide):
        r = item(r, key, text, shade=(i % 2 == 1))
    r += 1
    tip = ws.cell(row=r, column=2, value=ui(lang, 'filter_tip'))
    tip.font = Font(name=FONT, size=9, italic=True, color=RED_DARK)
    tip.alignment = Alignment(vertical='top', wrap_text=True)
    ws.merge_cells('B%d:%s%d' % (r, LAST, r))
    ws.row_dimensions[r].height = _height(ui(lang, 'filter_tip'), chars_per_line=115, line=12.5)
    r += 1

    if has_glossary:
        r += 1
        r = section(r, ui(lang, 'glossary_section'))
        ws.row_dimensions[r].height = 6
        r += 1
        for i, (key, text) in enumerate(GLOSSARY_GUIDE[lang]):
            r = item(r, key, text, shade=(i % 2 == 1))

    r += 2
    foot = ws.cell(row=r, column=2, value=ui(lang, 'footer', version=VERSION))
    foot.font = Font(name=FONT, size=9, color=RED_DARK)
    ws.merge_cells('B%d:%s%d' % (r, LAST, r))


# --- Catalogue --------------------------------------------------------------

def _sheet_catalog(ws, work, resolver, translator, compact=None, columns=None, per_rule=None,
                   lang='fr'):
    base_columns = COLUMNS[lang]
    columns = columns or []
    per_rule = per_rule or {}
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'D2' if columns else 'A2'
    wide = len(columns) > 12          # au-delà, en-têtes verticaux

    headers = list(base_columns) + [(title, 4.6 if wide else 11) for title, _k in columns]
    for idx, (name, width) in enumerate(headers, start=1):
        c = ws.cell(row=1, column=idx, value=name)
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=RED)
        rotated = wide and idx > len(base_columns)
        c.alignment = Alignment(vertical='center', horizontal='center',
                                wrap_text=not rotated,
                                textRotation=90 if rotated else 0)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(idx)].width = width
    if wide:
        longest = max((len(t) for t, _k in columns), default=10)
        ws.row_dimensions[1].height = max(108, min(250, longest * 6.2))
    else:
        ws.row_dimensions[1].height = 26

    row = 2
    for r in work:
        text, _ = describe(r, resolver, translator, compact)
        counter = int(r['Compteur ?']) == 1
        order = _as_int(r['Ordre'])
        values = [
            order,
            int(r['#']),
            clean_label(r['Libellé']),
            clean_label(r['Libellé court']),
            clean_code(r.get('Code')),
            text,
            ui(lang, 'yes') if counter else '',
            clean_label(r['Période'], inline=True) if counter else '—',
        ]
        marks = per_rule.get(r['_row'], set())
        values += ['X' if key in marks else '' for _title, key in columns]
        band = (row % 2 == 0)
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=value)
            c.font = Font(name=FONT, size=10, bold=(col == 3), color=INK)
            c.border = BORDER
            is_contract = col > len(base_columns)
            c.alignment = Alignment(vertical='center' if is_contract else 'top',
                                    wrap_text=(col in (3, 4, 6)),
                                    horizontal='center' if col in (1, 2, 7) or is_contract else 'left')
            if is_contract and value:
                c.font = Font(name=FONT, size=10, bold=True, color=RED)
            if counter and col in (7, 8):
                c.fill = PatternFill('solid', fgColor=RED_PALE)
                c.font = Font(name=FONT, size=10, bold=True, color=RED_DARK)
            elif band:
                c.fill = PatternFill('solid', fgColor=GREY)
        row += 1

    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(headers)), row - 1)


def _sheet_glossary(ws, details, usage, lang='fr'):
    """Lexique des types de jour : ce que le client voit dans ses plannings."""
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    columns = GLOSSARY_COLUMNS[lang]

    for idx, (name, width) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=idx, value=name)
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=RED)
        c.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 26

    # tri alphabétique : un lexique se consulte par le mot, pas par l'ordre d'écran
    details = sorted(details, key=lambda d: (d['public'] or d['court'] or '').lower())

    row = 2
    for item in details:
        nb = usage.get(item['id'], 0)
        duree = ''
        try:
            minutes = int(item['duree'])
            duree = ('%g h' % (minutes / 60)).replace('.', ',' if lang == 'fr' else '.')
        except (TypeError, ValueError):
            pass
        actif = ui(lang, 'yes') if item['actif'] else ui(lang, 'no')
        values = [item['public'], item['court'], item['id'], '', item['nature'],
                  duree, actif, nb]
        band = (row % 2 == 0)
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=value)
            c.font = Font(name=FONT, size=10, color=INK)
            c.border = BORDER
            c.alignment = Alignment(vertical='center', wrap_text=(col == 5),
                                    horizontal='center' if col in (3, 4, 6, 7, 8) else 'left')
            if col == 4 and item['couleur']:
                # la couleur réelle du type de jour, telle qu'elle apparaît au planning
                c.fill = PatternFill('solid', fgColor=item['couleur'])
                c.value = '#%d' % item['id']
                c.font = Font(name=FONT, size=9, bold=True, color=item['encre'])
            elif col == 8 and nb == 0:
                c.font = Font(name=FONT, size=10, color='9A9DA3', italic=True)
            elif band:
                c.fill = PatternFill('solid', fgColor=GREY)
        row += 1

    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(columns)), row - 1)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--input', required=True)
    ap.add_argument('--output', required=False)
    ap.add_argument('--daytypes')
    ap.add_argument('--client')
    ap.add_argument('--sample', help='liste de # de règles, pour la mise au point')
    ap.add_argument('--contrats', choices=['non', 'contrats'], default='contrats',
                    help='colonnes de filtre par type de contrat')
    ap.add_argument('--ref', action='append', default=[],
                    help="référentiel complémentaire : « famille=fichier.xlsx » ou « fichier.xlsx » "
                         "si le nom permet de deviner la famille. Répétable. Facultatif : sans lui "
                         "les références sortent entre chevrons.")
    ap.add_argument('--langue', choices=['fr', 'en'], default='fr',
                    help='langue de sortie du dossier — à demander à l\'utilisateur')
    ap.add_argument('--diagnostic', action='store_true',
                    help="n'écrit rien : affiche le contrôle préalable et les référentiels à réclamer")
    a = ap.parse_args()
    try:
        run(a)
    except ExportError as exc:
        print('ERREUR — %s' % exc, file=sys.stderr)
        sys.exit(2)


def run(a):
    rules = load_rules(a.input)
    days = load_daytypes(a.daytypes) if a.daytypes else {}
    referentials, report = load_referentials(a.ref)
    for line in report:
        print('Référentiel — %s' % line)
    if report:
        print()
    print(format_diagnosis(diagnose(rules, days, HANDLERS, referentials)))
    print()
    if a.diagnostic:
        return
    try:
        sample = [int(x) for x in a.sample.split(',') if x.strip()] if a.sample else None
    except ValueError:
        raise ExportError("--sample attend des numéros de règle séparés par des virgules : %s"
                          % a.sample)
    path = build(a.input, a.output, a.daytypes, a.client, sample, a.contrats, a.langue,
                 referentials)
    print('écrit :', path)


if __name__ == '__main__':
    main()
