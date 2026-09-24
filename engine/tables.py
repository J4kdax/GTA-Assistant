"""Lecture des classeurs sans pandas (depuis v2.0).

Le moteur tourne aussi dans le navigateur (Pyodide) : pandas et numpy y pèsent
une quinzaine de mégaoctets et plusieurs secondes de chargement pour un usage
qui se résume à lire des tableaux. openpyxl suffit.

Conventions, calquées sur ce que pandas faisait implicitement, pour que les
descriptions restent identiques :
  - une cellule vide vaut None ;
  - un nombre entier stocké en flottant (10.0) est ramené à l'entier (10) ;
  - les en-têtes sont des chaînes nettoyées de leurs espaces, un en-tête vide
    devient « Unnamed: N », un doublon reçoit le suffixe « .1 », « .2 »… ;
  - les lignes entièrement vides sont ignorées.
"""
import csv
import io
import os


class Sheet:
    """Une feuille : son nom, ses colonnes, ses lignes (dicts colonne -> valeur)."""

    def __init__(self, name, columns, rows):
        self.name = name
        self.columns = columns
        self.rows = rows

    def __len__(self):
        return len(self.rows)


class Table(list):
    """Liste de lignes (dicts) qui garde ses colonnes et des métadonnées.

    Remplace le DataFrame là où le moteur n'en utilisait qu'une fraction :
    `attrs` porte le format détecté et les anomalies de lecture.
    """

    def __init__(self, rows=(), columns=None, attrs=None):
        super().__init__(rows)
        self.columns = list(columns or [])
        self.attrs = dict(attrs or {})

    def derive(self, rows):
        """Même colonnes et métadonnées, autres lignes."""
        return Table(rows, self.columns, self.attrs)

    def column(self, name, default=None):
        return [row.get(name, default) for row in self]


def is_missing(value):
    return value is None or (isinstance(value, float) and value != value)


def _normalise(value):
    if isinstance(value, float):
        if value != value:
            return None
        if value.is_integer():
            return int(value)
    if isinstance(value, str) and value == '':
        return None
    return value


def _headers(raw):
    seen, out = {}, []
    for i, value in enumerate(raw):
        name = '' if value is None else str(value).strip()
        if not name:
            name = 'Unnamed: %d' % i
        if name in seen:
            seen[name] += 1
            name = '%s.%d' % (name, seen[name])
        else:
            seen[name] = 0
        out.append(name)
    return out


def _sheet_from_rows(name, rows, nrows=None):
    rows = iter(rows)
    header = next(rows, None)
    if header is None:
        return Sheet(name, [], [])
    columns = _headers(header)
    out = []
    for raw in rows:
        values = [_normalise(v) for v in raw]
        if all(v is None for v in values):
            continue
        values += [None] * (len(columns) - len(values))
        out.append(dict(zip(columns, values)))
        if nrows is not None and len(out) >= nrows:
            break
    return Sheet(name, columns, out)


def read_sheets(path, nrows=None):
    """Toutes les feuilles d'un classeur .xlsx, ou le contenu d'un .csv / .tsv.

    Lève FileNotFoundError si le fichier n'existe pas, ValueError s'il n'est
    pas lisible : l'appelant décide si c'est bloquant.
    """
    path = str(path)
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    if path.lower().endswith(('.csv', '.tsv')):
        with open(path, encoding='utf-8-sig', newline='') as fh:
            text = fh.read()
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel_tab if path.lower().endswith('.tsv') else csv.excel
        return [_sheet_from_rows(os.path.basename(path),
                                 csv.reader(io.StringIO(text), dialect), nrows)]
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError('%s: %s' % (type(exc).__name__, exc))
    try:
        return [_sheet_from_rows(ws.title, ws.iter_rows(values_only=True), nrows)
                for ws in wb.worksheets]
    finally:
        wb.close()
