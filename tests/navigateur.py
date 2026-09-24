"""Test de bout en bout de l'Atelier dans un vrai navigateur (Chromium, Playwright).

Sert le dépôt en local, dépose deux exports synthétiques, parcourt les cinq
onglets, télécharge le dossier et vérifie qu'il est identique, cellule par
cellule, à celui produit en ligne de commande. Échoue sur toute erreur de
console.

    pip install playwright && python -m playwright install chromium
    python3 tests/navigateur.py

Hors ligne ou derrière un proxy qui bloque jsDelivr : ATELIER_CDN_LOCAL=<dossier>
sert Pyodide et vis-network depuis des copies locales
(<dossier>/pyodide/… issu du paquet npm pyodide@0.26.2,
 <dossier>/vis-network/… issu du paquet npm vis-network@9.1.9).
"""
import asyncio
import functools
import http.server
import mimetypes
import os
import sys
import tempfile
import threading
import traceback

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(ROOT, 'engine'))
import robustesse as R          # noqa: E402
import build_xlsx               # noqa: E402

LOCAL = os.environ.get('ATELIER_CDN_LOCAL')


def serve():
    class Quiet(http.server.SimpleHTTPRequestHandler):
        def log_message(self, *args):
            pass
    handler = functools.partial(Quiet, directory=ROOT)
    httpd = http.server.ThreadingHTTPServer(('127.0.0.1', 0), handler)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd


def local_cdn(route):
    url = route.request.url.split('?')[0]
    if '/pyodide/v0.26.2/full/' in url:
        path = os.path.join(LOCAL, 'pyodide', url.split('/full/')[1])
    elif 'vis-network@9.1.9/standalone/umd/' in url:
        path = os.path.join(LOCAL, 'vis-network', url.rsplit('/', 1)[1])
    elif url.endswith('.css'):
        return route.fulfill(status=200, body='', content_type='text/css')
    else:
        return route.abort()
    if not os.path.exists(path):
        return route.fulfill(status=404, body='')
    ctype = 'application/wasm' if path.endswith('.wasm') else (
        mimetypes.guess_type(path)[0] or 'application/octet-stream')
    return route.fulfill(status=200, path=path, content_type=ctype)


# Mots français qui ne doivent plus être visibles une fois l'interface en anglais.
# Les libellés du client (règles, contrats, types de jour) sont exclus : ils
# restent dans leur langue d'origine par construction.
FRENCH_LEFT = r"""() => {
  const WORDS = /\b(règles?|Règles?|Aucune?|aucune?|Cliquez|Déposer|Déposez|dépendances?|Paramètres|Période|Compteur|Libellé|Accueil|Recherche|Fermer|Télécharger|Analyse|trouvé|chargé|référentiels?|affectations?|contrats?|niveau|formules?)\b/;
  const skip = el => el.closest('.lib, .short, .chip, .an-node .lib, .dt-chip, .affect, .params, .pkey, .pval, option, #file-list .fname, pre, .an-table td:nth-child(2), #detail h3, .rid .pill, .an-sub, .edge-rule, .opt-lib, code');
  const out = [];
  const walk = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
  let n;
  while ((n = walk.nextNode())) {
    const el = n.parentElement;
    if (!el || !n.textContent.trim()) continue;
    if (el.closest('[hidden], script, style')) continue;
    const r = el.getBoundingClientRect();
    if (!r.width || !r.height || getComputedStyle(el).visibility === 'hidden') continue;
    if (skip(el)) continue;
    const m = n.textContent.match(WORDS);
    if (m) out.push(n.textContent.trim().slice(0, 80));
  }
  return [...new Set(out)].slice(0, 12);
}"""


def cells(path):
    from openpyxl import load_workbook
    return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in load_workbook(path).worksheets]


async def run(tmp):
    from playwright.async_api import async_playwright
    rules = R.write(tmp, 'Contrats-GTA-Règles-test.xlsx', R.frame(R.FR))
    days = R.write(tmp, 'Contrats-GTA-Absences-présences-test.xlsx', sheets={'Export DIESE': R.table(**{
        '#': [7], 'Intitulé public': ['Travail'], 'Intitulé court public': ['TR'],
        'Absence': ['Temps de travail effectif'],
        'Couleur': ['<span style="background-color:#87CEFA;color:#000000">x</span>'],
        'Actif': [1], 'Durée défaut (min.)': [420]})})
    httpd = serve()
    base = 'http://127.0.0.1:%d/' % httpd.server_address[1]
    errors = []
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        ctx = await browser.new_context(viewport={'width': 1440, 'height': 900}, accept_downloads=True,
                                        locale='fr-FR')
        if LOCAL:
            await ctx.route('https://cdn.jsdelivr.net/**', local_cdn)
        page = await ctx.new_page()
        page.on('console', lambda m: errors.append(m.text) if m.type == 'error' else None)
        page.on('pageerror', lambda e: errors.append('pageerror: %s' % e))

        await page.goto(base + 'index.html')
        await page.wait_for_function("document.getElementById('engine-status').dataset.kind === 'ok'",
                                     timeout=180000)
        assert await page.is_disabled('nav.tabs [data-view=catalogue]'), 'onglets actifs sans données'
        await page.set_input_files('#files-input', [rules, days])
        await page.wait_for_function("!document.getElementById('env-summary').hidden", timeout=60000)
        files = await page.inner_text('#file-list')
        assert 'export des règles GTA' in files and 'export des types de jour' in files, files

        await page.click('nav.tabs [data-view=catalogue]')
        count = await page.text_content('#cat-count')
        assert count.startswith('3 / 3'), count

        await page.click('#cat-table a[data-rule="1"]')
        await page.wait_for_function("document.querySelector('#detail .rt-human')", timeout=10000)
        assert await page.get_attribute('#app', 'data-view') == 'explorer'

        await page.click('nav.tabs [data-view=audit]')
        assert 'Audit du paramétrage' in await page.inner_text('#audit-body')

        # ---- bascule en anglais : chaque onglet, aucun mot français visible
        await page.click('#lang-switch [data-lang=en]')
        assert await page.get_attribute('html', 'lang') == 'en'
        assert await page.get_attribute('#lang-switch [data-lang=en]', 'aria-checked') == 'true'
        tabs = await page.inner_text('nav.tabs')
        assert 'Home' in tabs and 'Client dossier' in tabs, tabs
        for view in ('accueil', 'catalogue', 'audit', 'dossier', 'explorer'):
            await page.click('nav.tabs [data-view=%s]' % view)
            await page.wait_for_timeout(150)
            french = await page.evaluate(FRENCH_LEFT)
            assert not french, 'français visible en anglais (%s) : %s' % (view, french)
        assert 'What the rule calculates' in await page.inner_text('#detail')
        await page.click('#btn-audit-side')
        await page.wait_for_timeout(150)
        french = await page.evaluate(FRENCH_LEFT)
        assert not french, 'français visible dans l\'audit du graphe : %s' % french
        await page.click('#an-close')
        await page.click('nav.tabs [data-view=catalogue]')
        assert 'What the rule calculates' in await page.inner_text('#cat-table thead')
        body = await page.inner_text('#cat-table tbody')
        assert 'Counts' in body or 'Adds up' in body or 'Calculates' in body, body[:300]

        # la langue est mémorisée
        await page.reload()
        await page.wait_for_function("document.getElementById('engine-status').dataset.kind === 'ok'",
                                     timeout=180000)
        assert await page.get_attribute('html', 'lang') == 'en'
        await page.set_input_files('#files-input', [rules, days])
        await page.wait_for_function("!document.getElementById('env-summary').hidden", timeout=60000)

        # le dossier suit la langue de l'interface
        await page.click('nav.tabs [data-view=dossier]')
        assert await page.input_value('#d-langue') == 'en'
        await page.fill('#d-client', 'Client test')
        async with page.expect_download(timeout=60000) as info:
            await page.click('#d-generate')
        download = await info.value
        web_en = os.path.join(tmp, 'web_en.xlsx')
        await download.save_as(web_en)
        assert download.suggested_filename.startswith('GTA-configuration-dossier_Client-test_')

        # retour au français
        await page.click('#lang-switch [data-lang=fr]')
        assert 'Accueil' in await page.inner_text('nav.tabs')
        assert await page.input_value('#d-langue') == 'fr'

        await page.click('#btn-theme')
        await page.click('#btn-theme')

        await page.click('nav.tabs [data-view=dossier]')
        async with page.expect_download(timeout=60000) as info:
            await page.click('#d-generate')
        download = await info.value
        web = os.path.join(tmp, 'web.xlsx')
        await download.save_as(web)
        assert download.suggested_filename.startswith('Dossier-parametrage-GTA_Client-test_')
        await browser.close()
    httpd.shutdown()

    cli = build_xlsx.build(rules, os.path.join(tmp, 'cli.xlsx'), days, 'Client test', None,
                           'contrats', 'fr')
    assert cells(web) == cells(cli), 'le dossier du navigateur diffère de la ligne de commande'
    cli_en = build_xlsx.build(rules, os.path.join(tmp, 'cli_en.xlsx'), days, 'Client test', None,
                              'contrats', 'en')
    assert cells(web_en) == cells(cli_en), 'le dossier anglais du navigateur diffère de la ligne de commande'
    assert not errors, errors


if __name__ == '__main__':
    with tempfile.TemporaryDirectory() as tmp:
        try:
            asyncio.run(run(tmp))
        except AssertionError as exc:
            traceback.print_exc()
            print("ÉCHEC — %s" % exc)
            sys.exit(1)
    print('navigateur : parcours complet ok (accueil, catalogue, explorer, audit, bascule FR/EN, '
          'mémorisation de la langue, thème, dossier FR et EN)')
