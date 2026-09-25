"""Tests de robustesse (depuis v1.9) — exports synthétiques construits à la volée.

Chaque cas fabrique un export #Dièse « abîmé » dans un dossier temporaire et
vérifie que le générateur :
  - produit un classeur exploitable quand l'anomalie est tolérable ;
  - refuse avec un message clair (ExportError) quand elle ne l'est pas ;
  - ne plante jamais sur une trace Python incompréhensible.

python3 tests/robustesse.py        → code de sortie non nul en cas d'échec.
Aucune donnée client : tout est synthétique.
"""
import os
import sys
import tempfile
import traceback

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'engine'))
import build_xlsx                                   # noqa: E402
from dossier_common import ExportError, load_rules  # noqa: E402

FR = ['#', 'Règle', 'Paramètres', 'Libellé', 'Libellé court', 'Code', 'Ordre',
      'Compteur ?', 'Période', 'Affectations', 'Affiché dièse ?']
EN = ['#', 'Rules', 'Parameters', 'Name', 'Short label', 'Code', 'Order', 'Counter?',
      'Period', 'Affectations', 'Display #DIESE?']
OLD = ['id', 'rule_type', 'parametres', 'libelle', 'libelle_court', 'code', 'ordre',
       'compteur', 'periode', 'affectations']

ROWS = [
    [1, '_039_STD_GTA_Formula', '_039_STD_GTA_Formula : rule2*2', 'Double <b>heures</b>',
     'Double<br>heures', None, 20, 0, None, 'TECH_16_FLEX - temps plein (#4)\nCDI adm (#7)', 1],
    [2, '_001_STD_GTA_Count_Hours',
     'Day types : 7\nMode* : Actual hours\nValue max* : 24\nInclude classical shifts* : Yes',
     'Heures', 'H', 'H01', 10, 1, 'Semaine', 'CDI adm (#7)', 0],
    [3, '_999_STD_GTA_Inconnu', 'Truc : 3', 'Règle inconnue', '', '???', 30, 0, None, None, 0],
]


class Frame:
    """Tableau minimal (en-têtes + lignes) pour fabriquer les exports de test."""

    def __init__(self, cols, rows):
        self.cols, self.rows = list(cols), [list(r) for r in rows]

    def drop(self, columns):
        keep = [i for i, c in enumerate(self.cols) if c not in columns]
        return Frame([self.cols[i] for i in keep], [[r[i] for i in keep] for r in self.rows])


def frame(cols, rows=ROWS):
    return Frame(cols, [r[:len(cols)] for r in rows])


def table(**columns):
    names = list(columns)
    return Frame(names, list(zip(*columns.values())))


def write(tmp, name, df=None, sheets=None):
    path = os.path.join(tmp, name)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet, d in (sheets or {'Export': df}).items():
        ws = wb.create_sheet(sheet)
        ws.append(d.cols)
        for row in d.rows:
            ws.append(row)
    wb.save(path)
    return path


def build(tmp, src, **kw):
    out = os.path.join(tmp, 'out_%s.xlsx' % os.path.basename(src))
    build_xlsx.build(src, out, contracts=kw.pop('contracts', 'contrats'), **kw)
    return load_workbook(out)


def catalog_rows(wb):
    ws = wb.worksheets[1]
    return [[c.value for c in row] for row in ws.iter_rows(min_row=2)]


CASES = []


def case(fn):
    CASES.append(fn)
    return fn


# --- ce qui doit passer ------------------------------------------------------

@case
def nominal_trois_habillages(tmp):
    for cols in (FR, EN, OLD):
        wb = build(tmp, write(tmp, 'r_%d.xlsx' % len(cols), frame(cols)))
        rows = catalog_rows(wb)
        assert [r[1] for r in rows] == [2, 1, 3], rows       # tri sur l'ordre
        assert all('<' not in str(v) for r in rows for v in r[2:4]), 'HTML resté dans un libellé'


@case
def export_vide(tmp):
    wb = build(tmp, write(tmp, 'vide.xlsx', frame(FR, [])))
    assert catalog_rows(wb) == []


@case
def valeurs_vides_partout(tmp):
    rows = [[4, '_001_STD_GTA_Count_Hours', None, None, None, None, None, None, None, None, None]]
    rows = catalog_rows(build(tmp, write(tmp, 'nan.xlsx', frame(FR, ROWS + rows))))
    assert len(rows) == 4


@case
def compteur_et_ordre_non_numeriques(tmp):
    rows = [r[:] for r in ROWS]
    rows[0][7], rows[1][7] = 'Oui', 'Yes'
    rows[2][6] = 'abc'
    wb = build(tmp, write(tmp, 'txt.xlsx', frame(FR, rows)))
    counters = {r[1]: r[6] for r in catalog_rows(wb)}
    assert counters[1] == 'Oui' and counters[2] == 'Oui', counters


@case
def lignes_sans_identifiant(tmp):
    rows = ROWS + [[None, '_001_STD_GTA_Count_Hours', '', 'ligne vide', '', '', 5, 0, '', '', 0],
                   ['abc', '_001_STD_GTA_Count_Hours', '', 'id illisible', '', '', 5, 0, '', '', 0]]
    rows = catalog_rows(build(tmp, write(tmp, 'noid.xlsx', frame(FR, rows))))
    assert sorted(r[1] for r in rows) == [1, 2, 3], rows


@case
def identifiants_en_double(tmp):
    rows = catalog_rows(build(tmp, write(tmp, 'dup.xlsx', frame(FR, ROWS + [ROWS[0]]))))
    assert len(rows) == 3, 'un doublon exact doit être dédoublonné'


@case
def feuille_regles_pas_en_premier(tmp):
    src = write(tmp, 'multi.xlsx', sheets={'Notes': table(x=[1]), 'Export': frame(FR)})
    assert len(catalog_rows(build(tmp, src))) == 3


@case
def formule_qui_fait_planter_le_traducteur(tmp):
    """Une exception dans le traducteur ne doit coûter qu'une ligne, pas le dossier."""
    import humanizer
    original = humanizer.FormulaTranslator.translate

    def boom(self, body):
        if 'BOOM' in body:
            raise RuntimeError('panne simulée')
        return original(self, body)
    humanizer.FormulaTranslator.translate = boom
    try:
        rows = ROWS + [[5, '_039_STD_GTA_Formula', '_039_STD_GTA_Formula : rule1+BOOM',
                        'Cassée', '', '', 40, 0, '', '', 0]]
        rows = {r[1]: r[5] for r in catalog_rows(build(tmp, write(tmp, 'boom.xlsx',
                                                                   frame(FR, rows))))}
    finally:
        humanizer.FormulaTranslator.translate = original
    assert len(rows) == 4 and 'panne simulée' in rows[5], rows[5]


@case
def types_de_jour_illisibles(tmp):
    bad = write(tmp, 'daytypes.xlsx', table(foo=[1], bar=[2]))
    wb = build(tmp, write(tmp, 'r.xlsx', frame(FR)), daytypes_path=bad)
    assert len(wb.worksheets) == 2


@case
def echantillon_hors_fichier(tmp):
    rows = catalog_rows(build(tmp, write(tmp, 's.xlsx', frame(FR)), sample=[2, 999]))
    assert [r[1] for r in rows] == [2]


@case
def sortie_par_defaut(tmp):
    src = write(tmp, 'Contrats-GTA-Règles-test.xlsx', frame(FR))
    out = build_xlsx.build(src, None)
    assert os.path.exists(out) and out.endswith('.xlsx'), out


@case
def balises_html_dans_les_libelles(tmp):
    """Toute balise HTML est retirée des libellés, son texte gardé ; les vraies
    comparaisons (« 00h < H < 08h ») ne sont pas prises pour des balises."""
    rows = [r[:] for r in ROWS]
    rows[0][3] = '<font color="red">Double</font> heures <sup>2</sup>&nbsp;x'
    rows[0][4] = '<small>GEN</small><br>Credit<br>Jour Maire'
    rows[1][4] = '00h < H < 08h [08]'
    rows[2][4] = '<SMALL>GEN</SMALL> < 35 h'
    got = {r[1]: (r[2], r[3]) for r in catalog_rows(build(tmp, write(tmp, 'tags.xlsx', frame(FR, rows))))}
    assert got[1] == ('Double heures 2 x', 'GEN\nCredit\nJour Maire'), got[1]
    assert got[2][1] == '00h < H < 08h [08]', got[2]
    assert got[3][1] == 'GEN < 35 h', got[3]


@case
def repli_sans_handler(tmp):
    rows = {r[1]: r[5] for r in catalog_rows(build(tmp, write(tmp, 'u.xlsx', frame(FR))))}
    assert rows[3], 'la règle sans handler doit garder une description de repli'


@case
def lisez_moi_aligne_sur_les_colonnes(tmp):
    """Le mode d'emploi du Lisez-moi suit exactement les en-têtes du catalogue."""
    for lang in ('fr', 'en'):
        guide = [k for k, _t in build_xlsx.COLUMN_GUIDE[lang]]
        assert guide == [n for n, _w in build_xlsx.COLUMNS[lang]], (lang, guide)
        assert len(build_xlsx.GLOSSARY_GUIDE[lang]) == len(build_xlsx.GLOSSARY_COLUMNS[lang]) - 1
    wb = build(tmp, write(tmp, 'lm.xlsx', frame(FR)), lang='en')
    text = ' '.join(str(c.value) for row in wb.worksheets[0].iter_rows() for c in row if c.value)
    assert 'CATALOGUE COLUMNS' in text and 'Short label' in text and 'TECH_16' in text


# --- ce qui doit être refusé proprement ------------------------------------

def refused(fn, *a, **kw):
    try:
        fn(*a, **kw)
    except ExportError as exc:
        assert str(exc) and 'Traceback' not in str(exc)
        return str(exc)
    raise AssertionError('aurait dû être refusé')


@case
def mauvais_fichier(tmp):
    src = write(tmp, 'daytypes_as_rules.xlsx',
                table(**{'#': [1], 'Intitulé public': ['Repos']}))
    msg = refused(load_rules, src)
    assert 'Règles GTA' in msg, msg


@case
def colonne_obligatoire_absente(tmp):
    msg = refused(load_rules, write(tmp, 'noaff.xlsx', frame(FR).drop(columns=['Affectations'])))
    assert 'Affectations' in msg, msg


@case
def fichier_absent_ou_pas_excel(tmp):
    refused(load_rules, os.path.join(tmp, 'nexiste_pas.xlsx'))
    txt = os.path.join(tmp, 'pas_excel.xlsx')
    open(txt, 'w').write('ceci n est pas un classeur')
    refused(load_rules, txt)


if __name__ == '__main__':
    failures = 0
    with tempfile.TemporaryDirectory() as tmp:
        for fn in CASES:
            try:
                fn(tmp)
                print('ok     %s' % fn.__name__)
            except Exception as exc:
                failures += 1
                print('ÉCHEC  %s — %s: %s' % (fn.__name__, type(exc).__name__, exc))
                if '-v' in sys.argv:
                    traceback.print_exc()
    print('%d cas, %d échec(s)' % (len(CASES), failures))
    sys.exit(1 if failures else 0)
